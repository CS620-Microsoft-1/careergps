from __future__ import annotations

import os
import subprocess
import sys
import tempfile
import io
import zipfile
from collections.abc import Callable
from pathlib import Path
from typing import Any

import pymupdf
from docx import Document

OCR_TEXT_THRESHOLD = 40

# Called as on_step(step, status, detail) when a pipeline step starts or finishes,
# e.g. on_step("ocr", "started", {"pages": [2]}), so the UI can show real progress.
StepCallback = Callable[[str, str, dict[str, Any] | None], None]


class UnsupportedResumeError(ValueError):
    """The upload is not a real PDF or DOCX file."""


class InvalidResumeError(ValueError):
    """The upload has the right type but cannot be safely read."""


class OCRProcessingError(RuntimeError):
    """OCR could not process an image-only page."""


def detect_file_type(data: bytes) -> str:
    if data.startswith(b"%PDF-"):
        return ".pdf"
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            names = set(archive.namelist())
            if "[Content_Types].xml" in names and "word/document.xml" in names:
                return ".docx"
    except (zipfile.BadZipFile, OSError):
        pass
    raise UnsupportedResumeError("Upload a valid PDF or DOCX file.")
def _docx_text(path: Path) -> str:
    try:
        document = Document(path)
    except Exception as error:
        raise InvalidResumeError("This DOCX file cannot be opened. Try exporting it again.") from error

    parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells if cell.text.strip())
    return "\n".join(parts)


def _page_text(page: pymupdf.Page) -> str:
    return page.get_text("text").strip()


def _read_pdf(path: Path) -> list[str]:
    try:
        with pymupdf.open(path) as document:
            if document.page_count == 0:
                raise InvalidResumeError("This PDF has no pages.")
            return [_page_text(page) for page in document]
    except InvalidResumeError:
        raise
    except Exception as error:
        raise InvalidResumeError("This PDF cannot be opened. Try exporting it again.") from error


def _pdf_text(
    path: Path,
    threshold: int = OCR_TEXT_THRESHOLD,
    on_step: StepCallback | None = None,
) -> dict[str, Any]:
    direct = _read_pdf(path)
    needs_ocr = [index for index, text in enumerate(direct) if len(text) < threshold]
    if not needs_ocr:
        return {"parser": "PyMuPDF", "ocr_pages": [], "text": "\n\n".join(direct)}

    # API page numbers are one-based even though PyMuPDF indexes pages from zero.
    ocr_pages = [index + 1 for index in needs_ocr]
    if on_step:
        on_step("ocr", "started", {"pages": ocr_pages})

    environment = os.environ.copy()
    tesseract_directory = Path(r"C:\Program Files\Tesseract-OCR")
    if tesseract_directory.exists():
        environment["PATH"] = str(tesseract_directory) + os.pathsep + environment.get("PATH", "")

    with tempfile.TemporaryDirectory(prefix="careergps_ocr_") as temporary_directory:
        output = Path(temporary_directory) / "ocr.pdf"
        command = [
            sys.executable,
            "-m",
            "ocrmypdf",
            "--output-type",
            "pdf",
            "--mode",
            "skip",
            "--language",
            "eng",
            str(path),
            str(output),
        ]
        try:
            completed = subprocess.run(
                command,
                env=environment,
                capture_output=True,
                text=True,
                timeout=120,
                check=False,
            )
        except (OSError, subprocess.TimeoutExpired) as error:
            raise OCRProcessingError(
                "OCR could not read this PDF. Try a clearer scan or export the file again."
            ) from error
        if completed.returncode != 0:
            raise OCRProcessingError(
                "OCR could not read this PDF. Try a clearer scan or export the file again."
            )
        final = _read_pdf(output)

    if on_step:
        on_step("ocr", "done", {"pages": ocr_pages})
    return {
        "parser": "PyMuPDF + OCRmyPDF",
        "ocr_pages": ocr_pages,
        "text": "\n\n".join(final),
    }


def extract_resume(path: Path, suffix: str, on_step: StepCallback | None = None) -> dict[str, Any]:
    if suffix == ".pdf":
        result = _pdf_text(path, on_step=on_step)
    elif suffix == ".docx":
        result = {"parser": "python-docx", "ocr_pages": [], "text": _docx_text(path)}
    else:
        raise UnsupportedResumeError("Upload a valid PDF or DOCX file.")

    result["text"] = result["text"].replace("\r\n", "\n").replace("\r", "\n").strip()
    if not result["text"]:
        raise InvalidResumeError("No text could be extracted from this résumé.")
    result["characters"] = len(result["text"])
    return result
